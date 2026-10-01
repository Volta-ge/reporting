r"""Extracts the finance team's own account -> sub-category label (the "Mapping" column of each TB
tab) from D:\all\volta\Finance\Volta FS montly Aug'26.xlsx, for grouping statement-line drill-downs
in dashboard_template.html by the finance team's own category names instead of raw account code
(2026-10-01, per user instruction - the current per-account-code grouping, e.g. inside Delivery
Expenses, is noisy for lines with many small leaf codes like "Other Opex" (7 4 *), where the finance
team's own Mapping already consolidates ~17 codes into 6 readable buckets: Office/Transport&Storage/
Utilities/Marketing/Office Rent/Management Service).

Base = TB Aug'26. Per the user ("თუ მანდ არ აღამოჩნდა შესაბამისი ანგარიშის მეპინგი, მაშინ წინა
თვეებში ნახე"), any account TB Aug'26 doesn't have a USABLE label for (absent entirely, blank, or
"Doesn't take Part in FS" - the aggregate/header-account marker, never a real leaf-posting label)
falls back through TB Jul'26 -> TB Jun'26 -> TB May'26 -> TB Apr'26 -> TB Mar'26 -> TB Feb'26, in
that priority order - whichever tab first has a real label wins. All 7 tabs share the identical
layout (header row 7; Account number/Type/Balance Mapping/Account Name/.../Mapping at column 13),
verified before writing this script.

Writes fs_mapping.json: {code: {label, name, source}}. NOT part of refresh.py - re-run by hand (then
build_mapping_class.py's consumer, build_dashboard.py) if the finance team supplies a newer copy or
a later month's TB tab should become the new base.
"""
import json
import os

import openpyxl

XLSX = os.environ.get('VOLTA_FS_XLSX', r"D:\all\volta\Finance\Volta FS montly Aug'26.xlsx")
BASE = os.environ.get('VOLTA_FIN_DATA', 'D:/all/volta/Volta_Accounting')
OUT = os.path.join(BASE, 'fs_mapping.json')

TABS_IN_PRIORITY = ["TB Aug'26", "TB Jul'26", "TB Jun'26", "TB May'26", "TB Apr'26", "TB Mar'26", "TB Feb'26"]
HEADER_ROW = 7
NO_LABEL = ("Doesn't take Part in FS",)
# 2026-10-01: briefly excluded '8 2 90 2'/'7 4 18' here, then reverted (the file's own Mapping column
# is the source of truth as written - no judgment calls against it without being asked). The user then
# asked again, specifically and explicitly, for both: each moved to the Office Expenses (95) statement
# line (was Unpredictable Expenses/104 - see the Sheet override + dashboard_template.html's
# OPX_EXPL7/MISC_CODES/AVPL_SPEC[95/104]) and should show there under its own account name, not folded
# into every other code's generic "ოფისის ხარჯი" file label.
BAD_LABEL = {'8 2 90 2', '7 4 18'}


def norm_code(raw):
    """Same convention as extract_mapping.py's norm_code(): a flat leading digit block becomes
    digit1 digit2 rest-of-block, with any already-space-separated sub-segments appended as-is -
    matches the Oris native WIRING code format (e.g. "7490 12" -> "7 4 90 12")."""
    if raw is None:
        return None
    raw = str(raw).strip()
    if not raw or raw in ('0', 'None'):
        return None
    parts = raw.split(' ', 1)
    head = parts[0]
    rest = parts[1] if len(parts) > 1 else ''
    if not head.isdigit() or len(head) < 3:
        return None
    out = head[0] + ' ' + head[1] + ' ' + head[2:]
    if rest:
        out += ' ' + rest.strip()
    return out


wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)

result = {}
per_tab_counts = {}
for tab in TABS_IN_PRIORITY:
    if tab not in wb.sheetnames:
        continue
    ws = wb[tab]
    added = 0
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        code = norm_code(row[0])
        if code is None or code in result or code in BAD_LABEL:
            continue
        mapping, name = row[12], row[4]
        label = str(mapping).strip() if mapping not in (None, '', 0, '0') else ''
        if not label or label in NO_LABEL:
            continue
        result[code] = {'label': label, 'name': str(name).strip() if name else '', 'source': tab}
        added += 1
    per_tab_counts[tab] = added

os.makedirs(BASE, exist_ok=True)
json.dump(result, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
print('wrote', OUT, len(result), 'accounts;', per_tab_counts)
